"""
Tests for the Network Compiler's format parsers (core/knowledge/parsers/)
and their wiring into the ingestion pipeline
(core/knowledge/enterprise/pipelines.py) and the SourceType taxonomy
(core/knowledge/enterprise/knowledge_layer.py).

Uses FakeEmbedder (core/knowledge/rag/embedder.py) so these tests never
download a real embedding model or touch the platform's real ChromaDB
store — each test gets its own temp-dir RAGEngine/EnterpriseKnowledgeLayer.
"""
import csv
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core.knowledge.enterprise.knowledge_layer import (
    SOURCE_RANK, EnterpriseKnowledgeLayer, SourceType,
)
from core.knowledge.enterprise.pipelines import ingest_directory, ingest_file
from core.knowledge.parsers import extract_text, supported_extensions
from core.knowledge.rag.embedder import FakeEmbedder
from core.knowledge.rag.rag_engine import RAGEngine


def _layer(tmp_path, name="test"):
    rag = RAGEngine(embedder=FakeEmbedder(), persist_dir=str(tmp_path / "chroma"),
                    collection_name=name)
    return EnterpriseKnowledgeLayer(rag=rag)


# ── SourceType taxonomy sanity ───────────────────────────────────────────────

def test_new_source_types_exist_and_ranked():
    expected = {
        "release_notes", "bug_report", "design_guide", "whitepaper",
        "golden_config", "customer_doc", "internal_wiki", "yang_model",
    }
    values = {s.value for s in SourceType}
    assert expected <= values
    for v in expected:
        assert v in SOURCE_RANK
        assert 0.0 < SOURCE_RANK[v] <= 1.0


# ── Per-format parser unit tests ────────────────────────────────────────────

def test_pdf_extract(tmp_path):
    from reportlab.pdfgen import canvas

    path = str(tmp_path / "guide.pdf")
    c = canvas.Canvas(path)
    c.drawString(72, 720, "OSPF adjacency requires matching MTU across neighbors.")
    c.save()

    text = extract_text(path)
    assert text and "matching MTU" in text


def test_docx_extract(tmp_path):
    from docx import Document

    path = str(tmp_path / "guide.docx")
    doc = Document()
    doc.add_paragraph("BGP peers must agree on remote-as before establishing a session.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "interface"
    table.rows[0].cells[1].text = "Gi0/1"
    doc.save(path)

    text = extract_text(path)
    assert text and "remote-as" in text
    assert "interface" in text and "Gi0/1" in text


def test_json_extract(tmp_path):
    path = str(tmp_path / "config.json")
    with open(path, "w") as f:
        json.dump({"interfaces": {"Gi0/1": {"mtu": 1500, "vlan": 10}}}, f)

    text = extract_text(path)
    assert text and "interfaces.Gi0/1.mtu: 1500" in text


def test_yaml_extract(tmp_path):
    path = str(tmp_path / "openconfig.yaml")
    with open(path, "w") as f:
        f.write("interfaces:\n  Gi0/1:\n    mtu: 1500\n")

    text = extract_text(path)
    assert text and "interfaces.Gi0/1.mtu: 1500" in text


def test_xml_extract(tmp_path):
    path = str(tmp_path / "yang.xml")
    with open(path, "w") as f:
        f.write("<config><interface name='Gi0/1'><mtu>1500</mtu></interface></config>")

    text = extract_text(path)
    assert text and "1500" in text and "Gi0/1" in text


def test_csv_extract(tmp_path):
    path = str(tmp_path / "vlans.csv")
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["vlan", "name"])
        w.writerow(["10", "USERS"])

    text = extract_text(path)
    assert text and "vlan=10" in text and "name=USERS" in text


def test_html_extract(tmp_path):
    path = str(tmp_path / "page.html")
    with open(path, "w") as f:
        f.write("<html><body><script>ignoreme()</script>"
                "<p>HSRP requires matching group numbers on both routers.</p>"
                "</body></html>")

    text = extract_text(path)
    assert text and "HSRP requires matching group numbers" in text
    assert "ignoreme" not in text


def test_unsupported_extension_returns_none(tmp_path):
    path = str(tmp_path / "binary.exe")
    with open(path, "wb") as f:
        f.write(b"\x00\x01\x02")
    assert extract_text(path) is None


def test_supported_extensions_cover_all_formats():
    exts = supported_extensions()
    for e in (".pdf", ".docx", ".json", ".yaml", ".yml", ".xml", ".csv", ".html", ".htm"):
        assert e in exts


# ── Pipeline integration: mixed-format directory ingest + retrieval ─────────

def test_ingest_directory_mixed_formats_and_search(tmp_path):
    src = tmp_path / "docs"
    src.mkdir()

    (src / "runbook.md").write_text(
        "# OSPF troubleshooting\nCheck neighbor state with show ip ospf neighbor.")

    with open(src / "vlans.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["vlan", "name"])
        w.writerow(["20", "VOICE"])

    with open(src / "iface.json", "w") as f:
        json.dump({"interface": {"Gi0/2": {"mtu": 9000}}}, f)

    layer = _layer(tmp_path, name="mixed")
    summary = ingest_directory(str(src), SourceType.DESIGN_GUIDE, vendor="cisco", layer=layer)

    assert summary["ingested"] == 3
    assert summary["skipped"] == 0
    assert not summary["errors"]

    hits = layer.search("OSPF neighbor state", top_k=3)
    assert any("neighbor" in h.text.lower() for h in hits)

    hits = layer.search("VOICE vlan", top_k=3)
    assert any("voice" in h.text.lower() for h in hits)

    hits = layer.search("Gi0/2 mtu", top_k=3)
    assert any("9000" in h.text for h in hits)


def test_ingest_file_single_pdf(tmp_path):
    from reportlab.pdfgen import canvas

    path = str(tmp_path / "whitepaper.pdf")
    c = canvas.Canvas(path)
    c.drawString(72, 720, "VXLAN uses UDP port 4789 for encapsulated traffic.")
    c.save()

    layer = _layer(tmp_path, name="pdf_single")
    r = ingest_file(path, SourceType.WHITEPAPER, layer=layer)
    assert not r.get("skipped")
    assert r["chunks"] >= 1

    hits = layer.search("VXLAN encapsulation port", top_k=3)
    assert hits and "4789" in hits[0].text


def test_ingest_directory_skips_unsupported_and_empty(tmp_path):
    src = tmp_path / "mixed2"
    src.mkdir()
    (src / "notes.txt").write_text("Golden config baseline for access switches.")
    (src / "image.png").write_bytes(b"\x89PNG\r\n\x1a\n")

    layer = _layer(tmp_path, name="skip_test")
    summary = ingest_directory(str(src), SourceType.GOLDEN_CONFIG, layer=layer)

    assert summary["ingested"] == 1
    assert summary["skipped"] == 0  # .png isn't in the extension allowlist, so it's not attempted
